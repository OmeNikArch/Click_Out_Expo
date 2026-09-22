#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Общий шаблон страницы «участники по эшелонам A/B/C» и выгрузки в Excel. Используют exhibitors_parkzoo.py и exhibitors_buybrand.py.
Запись: hall, stand, name, country, city, site, profile[], descr, products[], hl[] (жёлтые чипы), brands[], tags[], prio, top, meta (строка), _phones, _emails."""
import json, pathlib
from collections import Counter

ROOT = pathlib.Path(__file__).parent.parent
CSS = (ROOT / "events" / "exhibitors_build.py").read_text(encoding="utf-8").split('CSS = """')[1].split('"""')[0]

def page(cfg, rows):
    pub = [{k: v for k, v in r.items() if not k.startswith("_")} for r in rows]
    c = Counter(r["prio"] for r in rows); nTop = sum(r["top"] for r in rows)
    data = json.dumps(pub, ensure_ascii=False).replace("</", "<\\/")
    f1 = "".join(f'<button class="chip" data-f="expo" data-v="{v}">{l}</button>' for v, l in cfg["filter1"])
    legend = "".join(f'<div class="lg"><b>{k}</b> — {v}</div>' for k, v in cfg["legend"])
    return f"""<!doctype html>
<html lang="ru">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="robots" content="noindex,nofollow">
<title>{cfg['title']}</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link href="https://fonts.googleapis.com/css2?family=Unbounded:wght@500;700&family=JetBrains+Mono:wght@400;600&display=swap" rel="stylesheet">
<style>{CSS}</style>
</head>
<body>
<div class="wrap">
  <div class="kicker"><div class="bar"></div><span>{cfg['kicker']}</span></div>
  <h1>{cfg['h1'].format(n=len(rows))}</h1>
  <p class="sub">{cfg['sub'].format(n=len(rows), a=c['A'], b=c['B'], c=c['C'], top=nTop)}
  Презентация: <a href="presenter.html">версия ведущего</a> · <a href="index.html">клиентская</a> · <a href="../events.html">все выставки</a>.</p>
  <div class="legend">{legend}</div>
  <div class="panel">
    <div class="row"><input class="q" id="q" placeholder="{cfg['placeholder']}"><span class="cnt" id="cnt"></span></div>
    <div class="row"><span class="lbl">{cfg['filter1_label']}</span><button class="chip on" data-f="expo" data-v="">все</button>{f1}</div>
    <div class="row"><span class="lbl">Эшелон</span>
      <button class="chip" data-f="prio" data-v="top">★ первыми</button><button class="chip on" data-f="prio" data-v="A">A</button><button class="chip" data-f="prio" data-v="B">B</button><button class="chip" data-f="prio" data-v="C">C</button></div>
    <div class="row" id="tags"><span class="lbl">Сегмент</span></div>
  </div>
  <div class="list" id="list"></div>
</div>
<script>
const DATA={data};
const HALL_LABEL={json.dumps(cfg.get('hall_label', ''), ensure_ascii=False)};
const P=new URLSearchParams(location.search);
const state={{expo:(P.get("hall")||""),prio:new Set([(P.get("prio")||"A")]),tags:new Set(),q:""}};
document.querySelectorAll('.chip[data-f="expo"]').forEach(x=>x.classList.toggle("on",x.dataset.v===state.expo));
document.querySelectorAll('.chip[data-f="prio"]').forEach(x=>x.classList.toggle("on",state.prio.has(x.dataset.v)));
const TAGS=[...new Set(DATA.flatMap(r=>r.tags))].filter(t=>t!=="иностр.").sort();
const tagsRow=document.getElementById("tags");
TAGS.forEach(t=>{{const b=document.createElement("button");b.className="chip";b.dataset.f="tag";b.dataset.v=t;b.textContent=t;tagsRow.appendChild(b);}});
function esc(s){{return (s||"").replace(/[&<>"]/g,c=>({{"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;"}}[c]));}}
function host(u){{return u.replace(/^https?:[/][/]/,"").replace(/^www[.]/,"").replace(/[/].*$/,"")}}
function render(){{
  const q=state.q.trim().toLowerCase();
  const rows=DATA.filter(r=>{{
    if(state.expo&&r.hall!==state.expo)return false;
    if(state.prio.size){{let ok=false;state.prio.forEach(p=>{{if(p==="top"?r.top:r.prio===p)ok=true;}});if(!ok)return false;}}
    if(state.tags.size){{let ok=false;state.tags.forEach(t=>{{if(r.tags.includes(t))ok=true;}});if(!ok)return false;}}
    if(q){{const h=(r.name+" "+r.city+" "+r.site+" "+r.descr+" "+r.products.join(" ")+" "+r.brands.join(" ")+" "+r.stand+" "+(r.meta||"")).toLowerCase();if(!h.includes(q))return false;}}
    return true;}});
  document.getElementById("cnt").textContent=rows.length+" из "+DATA.length;
  const L=document.getElementById("list");
  if(!rows.length){{L.innerHTML='<div class="empty">Ничего не найдено — снимите фильтры.</div>';return;}}
  L.innerHTML=rows.map(r=>`<div class="it ${{r.prio}}">
    <div class="st">${{esc(r.stand||"—")}}<small>${{esc(HALL_LABEL?HALL_LABEL+" "+r.hall:r.hall)}}</small></div>
    <div><div class="nm"><span class="pr ${{r.prio}}">${{r.prio}}</span>${{r.top?'<span class="star">★</span>':''}}${{esc(r.name)}}</div>
      <div class="meta">${{esc([r.country!=="Россия"?r.country:"",r.city,r.profile.join(", "),r.meta||""].filter(Boolean).join(" · "))}}${{r.site?` · <a href="${{esc(r.site.startsWith("http")?r.site:"http://"+r.site)}}" target="_blank" rel="noopener">${{esc(host(r.site))}}</a>`:" · сайта нет"}}</div>
      <div class="tg">${{r.tags.map(t=>`<span>${{esc(t)}}</span>`).join("")}}${{(r.hl||[]).map(t=>`<span style="border-color:var(--y);color:var(--y)">${{esc(t)}}</span>`).join("")}}</div>
      ${{r.descr?`<div class="ds">${{esc(r.descr)}}</div>`:""}}
      ${{r.products.length||r.brands.length?`<div class="rb">${{r.brands.length?"Бренды: "+esc(r.brands.join(", "))+". ":""}}${{r.products.length?esc(r.products.join(" · ")):""}}</div>`:""}}
    </div></div>`).join("");
}}
document.addEventListener("click",e=>{{
  const b=e.target.closest(".chip");
  if(b){{const f=b.dataset.f,v=b.dataset.v;
    if(f==="expo"){{state.expo=v;document.querySelectorAll('.chip[data-f="expo"]').forEach(x=>x.classList.toggle("on",x.dataset.v===v));}}
    else{{const set=f==="prio"?state.prio:state.tags; set.has(v)?set.delete(v):set.add(v); b.classList.toggle("on");}}
    render();return;}}
  const it=e.target.closest(".it"); if(it&&e.target.closest(".ds,.rb")) it.classList.toggle("open");
}});
document.getElementById("q").addEventListener("input",e=>{{state.q=e.target.value;render();}});
render();
</script>
</body>
</html>
"""

def xlsx(cfg, rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); wb.remove(wb.active)
    cols = ["Эшелон", "★", cfg["filter1_label"], "Стенд", "Компания", "Страна", "Город", "Сайт", "Профиль", "Сегменты", cfg.get("hl_label", "Отметки"), "Бренды", "Описание", "Телефоны", "E-mail"]
    def sheet(title, rs):
        ws = wb.create_sheet(title); ws.append(cols)
        for c in ws[1]: c.font = Font(bold=True, color="000000"); c.fill = PatternFill("solid", fgColor="FDD101")
        for r in rs:
            ws.append([r["prio"], "★" if r["top"] else "", r["hall"], r["stand"], r["name"], r["country"], r["city"], r["site"], ", ".join(r["profile"]),
                       ", ".join(r["tags"]), ", ".join(r.get("hl", [])) + ((" · " + r["meta"]) if r.get("meta") else ""), ", ".join(r["brands"]), r["descr"][:400],
                       "; ".join(r.get("_phones", [])[:3]), "; ".join(r.get("_emails", [])[:3])])
        for col, w in zip("ABCDEFGHIJKLMNO", (7, 4, 14, 12, 34, 10, 14, 26, 18, 30, 30, 24, 60, 24, 28)): ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    sheet("A — Click Out", [r for r in rows if r["prio"] == "A"]); sheet("B — B2B", [r for r in rows if r["prio"] == "B"])
    sheet("C — не наши", [r for r in rows if r["prio"] == "C"]); sheet("Все", rows)
    wb.save(path)

def emit(cfg, rows, out_dir, argv):
    import sys
    out_dir.mkdir(exist_ok=True)
    pub = [{k: v for k, v in r.items() if not k.startswith("_")} for r in rows]
    (out_dir / "exhibitors.json").write_text(json.dumps(pub, ensure_ascii=False, indent=1), encoding="utf-8")
    (out_dir / "exhibitors.html").write_text(page(cfg, rows), encoding="utf-8")
    if "--xlsx" in argv:
        p = pathlib.Path(argv[argv.index("--xlsx") + 1]); p.parent.mkdir(parents=True, exist_ok=True); xlsx(cfg, rows, p); print("xlsx →", p)
    print(f"OK {out_dir.name}/exhibitors.html:", len(rows), dict(Counter(r["prio"] for r in rows)), "top", sum(r["top"] for r in rows))
    print("tags:", Counter(t for r in rows for t in r["tags"]).most_common())
